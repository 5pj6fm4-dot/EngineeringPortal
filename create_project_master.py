
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ==================================================
# 設定輸出路徑
# ==================================================
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "Data"
OUTPUT_FILE = DATA_DIR / "Project_Master.xlsx"

DATA_DIR.mkdir(parents=True, exist_ok=True)


# ==================================================
# 客戶主檔
# CustomerID 必須唯一，新增客戶時不要重複使用 ID
# ==================================================
customers = [
    ("CUS001", "Nexteer", "Y", ""),
    ("CUS002", "ZF", "Y", ""),
    ("CUS003", "Bosch", "Y", ""),
    ("CUS004", "BIBO", "Y", ""),
    ("CUS005", "BRP-Helium", "Y", ""),
    ("CUS006", "Johnson Outdoors", "Y", ""),
    ("CUS007", "TKP", "Y", ""),
    ("CUS008", "JD", "Y", ""),
    ("CUS009", "BHAP", "Y", ""),
    ("CUS010", "無", "Y", "通用／未指定客戶"),
]


# ==================================================
# 機種主檔
# ModelID 必須唯一；CustomerID 必須對應客戶主檔
# ==================================================
models = []
model_counter = 1


def add_models(customer_id, customer_name, model_names, remark=""):
    global model_counter

    for model_name in model_names:
        model_id = f"MOD{model_counter:03d}"

        models.append((
            model_id,
            customer_id,
            customer_name,
            model_name,
            "Y",
            remark
        ))

        model_counter += 1


add_models("CUS001", "Nexteer", [
    "31xx-2", "C1xx-2", "VSS", "Fiasa", "mCEPS", "無"
])

add_models("CUS002", "ZF", [
    "ERC", "AKC", "IBC2", "EMB", "EPBr", "SCA",
    "C1+", "IBC3", "BD", "TRW", "NEP3", "無"
])

add_models("CUS003", "Bosch", [
    "SCU4", "IPB", "無"
])

add_models("CUS004", "BIBO", ["無"])
add_models("CUS005", "BRP-Helium", ["無"])
add_models("CUS006", "Johnson Outdoors", ["無"])

add_models("CUS007", "TKP", [
    "SbW", "ICA3", "BL2B3", "C4.36", "C5.3", "無"
])

add_models("CUS008", "JD", ["CCE2", "無"])
add_models("CUS009", "BHAP", ["HBS", "無"])

add_models(
    "CUS010",
    "無",
    ["無"],
    "通用／未指定機種"
)


# ==================================================
# 建立活頁簿
# ==================================================
wb = Workbook()

ws_customer = wb.active
ws_customer.title = "CustomerDB"

ws_model = wb.create_sheet("ModelDB")
ws_readme = wb.create_sheet("README")


# ==================================================
# CustomerDB
# ==================================================
ws_customer.append([
    "CustomerID",
    "CustomerName",
    "IsActive",
    "Remark"
])

for customer in customers:
    ws_customer.append(customer)


# ==================================================
# ModelDB
# ==================================================
ws_model.append([
    "ModelID",
    "CustomerID",
    "CustomerName",
    "ModelName",
    "IsActive",
    "Remark"
])

for model in models:
    ws_model.append(model)


# ==================================================
# README
# ==================================================
instructions = [
    ["Project_Master.xlsx 使用說明"],
    ["項目", "說明"],
    ["用途", "提供 EngineeringPortal 各工程工具共用的客戶與機種資料。"],
    ["CustomerDB", "管理客戶名稱、客戶 ID、啟用狀態與備註。"],
    ["ModelDB", "管理機種名稱，以及機種所屬客戶。"],
    ["新增客戶", "在 CustomerDB 表格最後新增一列，CustomerID 不可重複。"],
    ["新增機種", "在 ModelDB 表格最後新增一列，ModelID 不可重複，CustomerID 必須存在於 CustomerDB。"],
    ["IsActive", "Y 表示啟用；N 表示停用。停用資料仍保留，供歷史資料參照。"],
    ["ID 注意事項", "CustomerID 與 ModelID 是識別碼，不建議因顯示名稱變更而修改。"],
    ["網頁連接", "網頁應透過 server.py 讀取此檔案，不要依賴瀏覽器直接讀取本機 Excel。"],
    ["檔案覆寫", "重新執行本程式會重新建立檔案，請先備份已經人工修改的資料。"],
]

for row in instructions:
    ws_readme.append(row)


# ==================================================
# 格式設定
# ==================================================
header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

bottom_border = Border(
    bottom=Side(style="thin", color="D9E2F3")
)


def format_sheet(ws, table_name):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )
            cell.border = bottom_border

    ws.row_dimensions[1].height = 25

    for col_idx in range(1, ws.max_column + 1):
        values = [
            len(str(ws.cell(r, col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        ]

        width = min(max(max(values) + 3, 12), 32)
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    if ws.max_row >= 2:
        table = Table(
            displayName=table_name,
            ref=ws.dimensions
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        ws.add_table(table)


format_sheet(ws_customer, "tblCustomerDB")
format_sheet(ws_model, "tblModelDB")


# ==================================================
# 啟用狀態下拉選單：Y / N
# ==================================================
for ws, col in [
    (ws_customer, "C"),
    (ws_model, "E")
]:
    validation = DataValidation(
        type="list",
        formula1='"Y,N"',
        allow_blank=False
    )

    ws.add_data_validation(validation)
    validation.add(f"{col}2:{col}1000")


# ==================================================
# README 格式
# ==================================================
ws_readme.column_dimensions["A"].width = 24
ws_readme.column_dimensions["B"].width = 100
ws_readme.freeze_panes = "A2"

for cell in ws_readme[1]:
    cell.fill = header_fill
    cell.font = header_font

for row in ws_readme.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

ws_readme.row_dimensions[1].height = 25


# ==================================================
# 儲存檔案
# ==================================================
wb.save(OUTPUT_FILE)

print("Project_Master.xlsx 建立完成")
print(f"檔案位置：{OUTPUT_FILE}")
print(f"客戶筆數：{len(customers)}")
print(f"機種筆數：{len(models)}")
